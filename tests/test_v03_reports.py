"""v0.3: dashboards, standard reports, report builder, exports, saved reports, schedules and escalation."""

import io
from datetime import date, datetime, timedelta

import pytest
from django.core import mail
from django.http import QueryDict
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import AuditEvent, Centre, Notification, Parameter
from apps.members import services
from apps.members.models import Application
from apps.reports import charts, datasets, escalation, export, standard
from apps.reports.filters import period_from
from apps.reports.models import ReportSchedule, SavedReport
from apps.reports.scheduling import send_schedule
from apps.rules.models import RejectionReason
from apps.security.models import SecurityAlert

from .conftest import PASSWORD, make_user
from .test_counter_journey import complete, start

pytestmark = pytest.mark.django_db


@pytest.fixture(autouse=True)
def no_mfa(settings):
    settings.MFA_REQUIRED = False
    settings.MFA_REQUIRED_ROLES = []


@pytest.fixture
def supervisor(centre):
    return make_user("sup.here", "supervisor", "customer_care", centre=centre)


@pytest.fixture
def other_supervisor(reference):
    return make_user("sup.there", "supervisor", "customer_care", centre=Centre.objects.get(code="MBJ"))


@pytest.fixture
def activity(agent, taj, doctor, run_commits):
    """Two approvals and a rejection at the agent's centre."""
    a = complete(start(agent, run_commits, "100200301", "Marcia", "Campbell", date(1979, 3, 14)), agent, doctor, "hypertension")
    with run_commits():
        services.approve(a.pk, user=agent)
    b = complete(start(agent, run_commits, "100200303", "Winston", "Grant", date(1958, 11, 20)), agent, doctor, "diabetes")
    with run_commits():
        services.approve(b.pk, user=agent)
    c = start(agent, run_commits, "100200305", "Paulette", "Reid", date(1972, 9, 1))
    services.reject(c, user=agent, reason=RejectionReason.objects.get(code="incomplete_form"))
    return a, b, c


def login(client, user):
    assert client.login(username=user.username, password=PASSWORD)


def period(user, q="preset=30"):
    return period_from(QueryDict(q), user)[0]


# --- access ------------------------------------------------------------------------------------
def test_reports_need_a_reporting_role(client, agent, manager):
    login(client, agent)
    assert client.get("/reports/dashboard/").status_code == 403
    assert client.get("/reports/").status_code == 403
    client.logout()
    login(client, manager)
    assert client.get("/reports/dashboard/").status_code == 200
    assert b"Dashboard" in client.get("/").content  # menu link


def test_health_report_only_for_clinical_roles(client, auditor, manager):
    login(client, auditor)
    assert client.get("/reports/standard/conditions/").status_code == 403
    assert b"Members by medical condition" not in client.get("/reports/").content
    client.logout()
    login(client, manager)
    assert client.get("/reports/standard/conditions/").status_code == 200


# --- dashboard and standard reports ------------------------------------------------------------------
def test_dashboard_counts(client, manager, activity):
    login(client, manager)
    page = client.get("/reports/dashboard/?preset=7").content.decode()
    assert "Applications received" in page and "Approval rate 66.7%" in page
    partial = client.get("/reports/dashboard/?preset=7", HTTP_HX_REQUEST="true").content.decode()
    assert "<html" not in partial and "Updated" in partial  # HTMX refresh returns the body only


def test_every_standard_report_renders_and_exports(client, manager, activity):
    login(client, manager)
    for key in standard.CATALOGUE:
        assert client.get(f"/reports/standard/{key}/?preset=30").status_code == 200, key
    resp = client.get("/reports/standard/applications/?preset=30&format=xlsx")
    wb = load_workbook(io.BytesIO(resp.content))
    assert "Applications by centre" in wb.sheetnames and resp["Content-Disposition"].endswith('.xlsx"')
    assert AuditEvent.objects.filter(action="report.exported").count() == 1


def test_applications_report_figures(manager, activity, centre):
    report = standard.build("applications", period(manager), manager)
    by_centre = report.tables[0]
    row = next(r for r in by_centre.rows if r[0] == centre.name)
    assert row[1:7] == [3, 3, 0, 2, 1, 0] and row[7] == "66.7%"


def test_supervisor_held_to_own_centre(manager, other_supervisor, activity):
    assert standard.build("applications", period(manager), manager).tables[0].rows
    p = period(other_supervisor, "preset=30&centre=")  # trying "all centres"
    assert p.centre.code == "MBJ"
    assert standard.build("applications", p, other_supervisor).tables[0].rows == []
    headers, rows, _ = datasets.run({"dataset": "applications", "group_by": ["status"]}, other_supervisor)
    assert rows == []


def test_personal_tables_hidden_without_personal_role(manager, activity):
    report = standard.build("staff", period(manager), manager)
    assert report.tables and report.tables[0].personal


# --- builder ------------------------------------------------------------------------------------------
def test_builder_groups_and_measures(manager, activity):
    headers, rows, truncated = datasets.run({"dataset": "applications", "group_by": ["status"], "measure": "count"}, manager)
    assert headers == ["Status", "Number of records"] and dict(rows) == {"Approved": 2, "Rejected": 1} and not truncated
    headers, rows, _ = datasets.run({"dataset": "applications", "group_by": ["channel"], "measure": "avg:turnaround"}, manager)
    assert headers[-1].startswith("Average turnaround") and isinstance(rows[0][1], float)


def test_builder_filters_and_validation(manager, activity):
    _, rows, _ = datasets.run({"dataset": "applications", "group_by": ["status"], "filters": [{"field": "status", "op": "eq", "value": "approved"}]}, manager)
    assert rows == [["Approved", 2]]
    with pytest.raises(datasets.ReportError, match="choose one of the listed values"):
        datasets.run({"dataset": "applications", "group_by": ["status"], "filters": [{"field": "status", "op": "eq", "value": "nonsense"}]}, manager)
    with pytest.raises(datasets.ReportError, match="YYYY-MM-DD"):
        datasets.run({"dataset": "applications", "group_by": ["status"], "filters": [{"field": "received_day", "op": "gte", "value": "yesterday"}]}, manager)
    with pytest.raises(datasets.ReportError, match="does not offer"):
        datasets.run({"dataset": "applications", "group_by": ["password"]}, manager)


def test_builder_refuses_health_fields_to_auditor(auditor, activity):
    with pytest.raises(datasets.ReportError, match="cannot see Medical condition"):
        datasets.run({"dataset": "members", "group_by": ["condition"]}, auditor)
    assert "condition" not in [f.key for f in datasets.get("members").visible_fields(auditor)]


def test_trn_always_masked(client, manager, activity):
    headers, rows, _ = datasets.run({"dataset": "members", "columns": ["member_number", "trn"]}, manager)
    assert all(r[1].startswith("•••-•••-") and len(r[1]) == 11 for r in rows)
    login(client, manager)
    body = client.get("/reports/builder/?dataset=members&col=trn&col=surname&format=csv").content.decode("utf-8-sig")
    assert "100200301" not in body and "•••-•••-301" in body
    assert AuditEvent.objects.filter(action="report.exported", summary__contains="personal columns").exists()


def test_builder_page_runs_and_saves(client, manager, activity):
    login(client, manager)
    page = client.get("/reports/builder/?dataset=applications&group=status&measure=count&run=1").content.decode()
    assert "Approved" in page and "2 rows" in page
    resp = client.post("/reports/builder/save/", {"dataset": "applications", "group": "status", "measure": "count", "name": "By status"})
    saved = SavedReport.objects.get(name="By status")
    assert resp.status_code == 302 and saved.aggregate_only and saved.definition["group_by"] == ["status"]
    assert "Approved" in client.get(f"/reports/saved/{saved.pk}/").content.decode()


def test_private_reports_stay_private(client, manager, activity, centre):
    other = make_user("mgr.other", "manager", centre=centre)
    mine = SavedReport.objects.create(name="Mine", definition={"dataset": "applications", "group_by": ["status"]}, owner=manager)
    login(client, other)
    assert client.get(f"/reports/saved/{mine.pk}/").status_code == 404
    mine.shared = True
    mine.save()
    assert client.get(f"/reports/saved/{mine.pk}/").status_code == 200
    assert client.post(f"/reports/saved/{mine.pk}/delete/").status_code == 404  # only the owner deletes


# --- exports --------------------------------------------------------------------------------------------
def test_exports_cannot_carry_formulas():
    data = export.xlsx_bytes([("Sheet", ["Name"], [['=HYPERLINK("http://evil")'], ["+SUM(A1)"], [42]])], title="t", meta=[])
    ws = load_workbook(io.BytesIO(data))["Sheet"]
    values = [c.value for c in ws["A"] if c.value is not None]
    assert '\'=HYPERLINK("http://evil")' in values and "'+SUM(A1)" in values and 42 in values
    csv_text = export.csv_bytes([("", ["Name"], [["=1+1"]])], title="t", meta=[]).decode("utf-8-sig")
    assert "'=1+1" in csv_text


def test_export_rate_limit(client, manager, activity):
    Parameter.objects.update_or_create(key="limit-report-exports-per-hour", defaults={"value": 5})
    login(client, manager)
    for _ in range(5):
        assert client.get("/reports/standard/rejections/?format=csv")["Content-Type"].startswith("text/csv")
    assert client.get("/reports/standard/rejections/?format=csv").status_code == 302
    assert SecurityAlert.objects.filter(kind="rate_limit", user=manager).exists()


# --- schedules and escalation ------------------------------------------------------------------------------
def test_only_count_reports_can_be_scheduled(client, manager, activity):
    rows_report = SavedReport.objects.create(name="List", definition={"dataset": "members", "columns": ["member_number"]}, owner=manager)
    login(client, manager)
    resp = client.get(f"/reports/saved/{rows_report.pk}/schedule/")
    assert resp.status_code == 302 and not rows_report.aggregate_only


def test_schedule_sends_role_scoped_excel_without_personal_tables(client, manager, agent, activity, centre):
    report = SavedReport.objects.create(name="Staff weekly", kind="standard", standard_key="staff", definition={}, owner=manager)
    login(client, manager)
    resp = client.post(f"/reports/saved/{report.pk}/schedule/", {"frequency": "weekly", "hour": 7, "period_days": 7, "recipients": [manager.pk]})
    assert resp.status_code == 302
    sched = ReportSchedule.objects.get(report=report)
    assert timezone.localtime(sched.next_run_at).weekday() == 0 and timezone.localtime(sched.next_run_at).hour == 7
    sched.recipients.add(agent)  # Customer Care has no reporting role: skipped
    result = send_schedule(sched)
    assert "sent to 1" in result and "1 skipped" in result
    assert len(mail.outbox) == 1 and mail.outbox[0].to == [manager.email]
    name, content, mime = mail.outbox[0].attachments[0]
    assert name.endswith(".xlsx") and mime == export.XLSX
    assert "Work by staff member" not in load_workbook(io.BytesIO(content)).sheetnames  # names staff: never e-mailed
    assert Notification.objects.filter(template="scheduled-report", recipient_user=manager).exists()


def test_builder_schedule_applies_the_period(manager, activity):
    report = SavedReport.objects.create(name="By status", definition={"dataset": "applications", "group_by": ["status"]}, owner=manager)
    sched = ReportSchedule.objects.create(report=report, frequency="daily", hour=6, period_days=7, created_by=manager)
    sched.recipients.add(manager)
    Application.objects.update(received_at=timezone.now() - timedelta(days=2))
    send_schedule(sched)
    ws = load_workbook(io.BytesIO(mail.outbox[0].attachments[0][1])).active
    assert any(row and row[0] == "Approved" and row[1] == 2 for row in ws.iter_rows(values_only=True))


def test_next_run_calculation(manager):
    report = SavedReport.objects.create(name="r", definition={"dataset": "applications", "group_by": ["status"]}, owner=manager)
    sched = ReportSchedule(report=report, frequency="monthly", hour=8, created_by=manager)
    nxt = timezone.localtime(sched.compute_next_run(timezone.make_aware(datetime(2026, 10, 30, 12, 0))))
    assert (nxt.month, nxt.day, nxt.hour) == (11, 1, 8)


def test_overdue_applications_escalate_once_a_day(agent, supervisor, taj, run_commits):
    app = start(agent, run_commits, "100200301", "Marcia", "Campbell", date(1979, 3, 14))
    Application.objects.filter(pk=app.pk).update(received_at=timezone.now() - timedelta(hours=30))
    assert escalation.run()["applications"] == 1
    note = Notification.objects.get(recipient_user=supervisor, template="overdue-applications")
    assert app.tracking_number in note.body
    assert escalation.run()["applications"] == 0  # no repeat within the day


# --- safety of the drawing and filters ------------------------------------------------------------------------
def test_chart_labels_are_escaped():
    svg = str(charts.bars([("<script>alert(1)</script>", 3)]))
    assert "<script>" not in svg and "&lt;script&gt;" in svg


def test_period_bounds(manager):
    p = period(manager, "date_from=2026-10-10&date_to=2026-10-01")
    assert p.start_date == date(2026, 10, 1) and p.end_date == date(2026, 10, 10)
    p = period(manager, "date_from=2020-01-01&date_to=2026-10-01")
    assert (p.end_date - p.start_date).days == 731


def test_run_scheduled_jobs_command(manager, activity):
    from django.core.management import call_command

    report = SavedReport.objects.create(name="Daily", kind="standard", standard_key="rejections", definition={}, owner=manager)
    sched = ReportSchedule.objects.create(report=report, frequency="daily", hour=6, created_by=manager, next_run_at=timezone.now() - timedelta(minutes=1))
    sched.recipients.add(manager)
    out = io.StringIO()
    call_command("run_scheduled_jobs", stdout=out)
    sched.refresh_from_db()
    assert "1 schedule(s) sent" in out.getvalue() and len(mail.outbox) == 1 and sched.next_run_at > timezone.now()
