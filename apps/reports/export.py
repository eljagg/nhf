"""CSV and Excel exports. Text that a spreadsheet would run as a formula is neutralised in both formats."""

import csv
import io

from django.http import HttpResponse
from django.utils import timezone

from apps.core import csvsafe

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _safe(value):
    if isinstance(value, int | float) and not isinstance(value, bool):
        return value
    return csvsafe.cell(value)


def csv_bytes(sections, *, title, meta):
    """sections = [(heading, headers, rows)]."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([csvsafe.cell(title)])
    for line in meta:
        w.writerow([csvsafe.cell(line)])
    for heading, headers, rows in sections:
        w.writerow([])
        if heading:
            w.writerow([csvsafe.cell(heading)])
        w.writerow(csvsafe.row(headers))
        for r in rows:
            w.writerow([_safe(v) for v in r])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM so Excel opens the dashes and bullets correctly


def xlsx_bytes(sections, *, title, meta):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used = set()
    for i, (heading, headers, rows) in enumerate(sections or [("", [], [])]):
        name = "".join(ch for ch in (heading or title) if ch not in "[]:*?/\\")[:28] or f"Sheet {i + 1}"
        while name in used:
            name = f"{name[:25]} {i + 1}"
        used.add(name)
        ws = wb.create_sheet(name)
        ws.append([csvsafe.cell(title)])
        ws["A1"].font = Font(bold=True, size=13)
        for line in meta:
            ws.append([csvsafe.cell(line)])
        if heading:
            ws.append([csvsafe.cell(heading)])
            ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.append([])
        ws.append(csvsafe.row(headers))
        header_row = ws.max_row
        for c in range(1, len(headers) + 1):
            cell = ws.cell(row=header_row, column=c)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="137062")
            cell.alignment = Alignment(vertical="center")
        for r in rows:
            ws.append([_safe(v) for v in r])
            for c in range(1, len(r) + 1):  # a value that still begins with "=" is stored as text, never as a formula
                cell = ws.cell(row=ws.max_row, column=c)
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
        for c, h in enumerate(headers, start=1):
            width = max([len(str(h))] + [len(str(r[c - 1])) for r in rows[:500] if len(r) >= c] + [8])
            ws.column_dimensions[get_column_letter(c)].width = min(width + 2, 60)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def response(sections, *, title, meta, fmt, stem):
    stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
    name = csvsafe.filename(f"{stem}-{stamp}")
    if fmt == "xlsx":
        resp = HttpResponse(xlsx_bytes(sections, title=title, meta=meta), content_type=XLSX)
        resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    else:
        resp = HttpResponse(csv_bytes(sections, title=title, meta=meta), content_type="text/csv; charset=utf-8")
        resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
    resp["Cache-Control"] = "no-store"
    return resp
